"""Exercise template routing and the actual Qt worker/save pipeline."""

import os
import time
from pathlib import Path

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

import openpyxl
import pytest
from PySide6.QtWidgets import QApplication, QMessageBox, QFileDialog

from tests.test_simple_report import sources, template
from src.gui.main_window import MainWindow


@pytest.mark.parametrize('include_discounts', [True, False])
def test_simple_report_upload_process_and_repeat(tmp_path, monkeypatch, sources, template, shift_config, include_discounts):
    app = QApplication.instance() or QApplication([])
    errors = []
    monkeypatch.setattr(QMessageBox, 'warning', lambda *args: errors.append(args[2]))
    monkeypatch.setattr(QMessageBox, 'critical', lambda *args: errors.append(args[2]))
    window = MainWindow()
    window.shift_config = shift_config
    for name in ('output_dir', 'table_dir', 'image_dir'):
        folder = tmp_path / name
        folder.mkdir()
        setattr(window, name, folder)
    original = template.read_bytes()
    window._handle_table_upload(template, 'simple_output_table')
    assert window.simple_report
    assert window.simple_output_table_path != template
    assert not window.simple_options.isHidden()
    assert window.simple_divisor.value() == window.simple_count_divisor.value() == 3
    window.simple_closing_balance.setText('608377.84')
    for category, path in sources.items():
        if category != '油品优惠' or include_discounts:
            window._handle_table_upload(path, category)
    assert window._validate_required_files()  # No images or API key needed.

    for _ in range(2):
        window._process_all_files()
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            app.processEvents()
            status = window.processing_status.get('simple_output_table', '')
            if status == '已填写，待核对保留项' or errors:
                break
            time.sleep(0.01)
        assert not errors
        assert status == '已填写，待核对保留项'
        assert not window.simple_pending_updates
        wb = openpyxl.load_workbook(window.simple_output_table_path)
        assert wb['Sheet1']['D11'].value == 164
        assert wb['Sheet1']['G7'].value == 202792.61
        assert wb['Sheet1']['B3'].value == 10
        assert wb['Sheet1']['E3'].value == 0
        assert wb['Sheet1']['E3'].number_format == '0'
        assert wb['Sheet2']['B11'].value is None
        assert wb['Sheet2']['C11'].value is None
        if not include_discounts:
            assert window.processing_status.get('油品优惠') != '完成'
        wb.close()
        app.processEvents()

    # The prominent output button saves the processed copy, not the uploaded template.
    exported = tmp_path / 'exported.xlsx'
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *args: (str(exported), ''))
    monkeypatch.setattr(QMessageBox, 'information', lambda *args: None)
    assert window.template_export_buttons['simple_output_table'].text() == '导出简表'
    window.template_export_buttons['simple_output_table'].click()
    assert exported.read_bytes() == window.simple_output_table_path.read_bytes()
    assert template.read_bytes() == original
    # Re-uploading the already-managed copy must not delete its own source.
    managed_path = window.simple_output_table_path
    window._handle_table_upload(managed_path, 'simple_output_table')
    assert managed_path.exists()
    assert not errors
    window.close()
    app.processEvents()


def test_legacy_processes_without_guotong(tmp_path, monkeypatch, template, sources, shift_config):
    app = QApplication.instance() or QApplication([])
    errors = []
    monkeypatch.setattr(QMessageBox, 'warning', lambda *args: errors.append(args[2]))
    monkeypatch.setattr(QMessageBox, 'critical', lambda *args: errors.append(args[2]))
    window = MainWindow()
    window.shift_config = shift_config
    for name in ('output_dir', 'table_dir', 'image_dir'):
        folder = tmp_path / name
        folder.mkdir()
        setattr(window, name, folder)
    legacy = tmp_path / 'legacy.xlsx'
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = '调价前'
    for row, section, machine in [(3, 'A', 3), (4, 'B', 9), (5, 'C', 11)]:
        sheet.cell(row, 1, section)
        sheet.cell(row, 2, machine)
    sheet['H92'] = 123  # No screenshot: preserve existing Guotong values.
    sheet['H93'] = 456
    discounts = workbook.create_sheet('油品优惠明细 2')
    discounts['B2'] = shift_config.date.day
    discounts['J2'] = 0
    workbook.save(legacy)
    window._handle_table_upload(legacy, 'output_table')
    assert not window.simple_report
    assert window.simple_options.isHidden()
    assert not window._validate_required_files()  # Still needs the source tables.
    assert not window.required_rows
    assert {'国通1', '国通2'}.issubset(window.optional_rows)
    for category, path in sources.items():
        window._handle_table_upload(path, category)
    assert window._validate_required_files()
    window._process_all_files()
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        app.processEvents()
        if window.processing_status.get('output_table') == '完成' or errors:
            break
        time.sleep(0.01)
    assert not errors
    assert window.processing_status['output_table'] == '完成'
    saved = openpyxl.load_workbook(window.output_table_path)
    assert saved['调价前']['D3'].value == 10
    assert saved['调价前']['H76'].value == 78
    assert saved['调价前']['H92'].value == 123
    assert saved['调价前']['H93'].value == 456
    saved.close()
    window.close()
    app.processEvents()


def test_export_waits_for_processing(tmp_path, monkeypatch, template):
    app = QApplication.instance() or QApplication([])
    messages = []
    monkeypatch.setattr(QMessageBox, 'warning', lambda *args: messages.append(args[2]))
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *args: pytest.fail('Must not export while processing'))
    window = MainWindow()
    window.output_dir = tmp_path / 'output'
    window.output_dir.mkdir()
    window._handle_table_upload(template, 'simple_output_table')
    window.files_to_process = 1
    window.template_export_buttons['simple_output_table'].click()
    window.files_to_process = 0
    window.processing_status['simple_output_table'] = '更新中'
    window.template_export_buttons['simple_output_table'].click()
    assert messages == ['表格正在处理，请完成后再输出'] * 2
    window.processing_status['simple_output_table'] = '已上传'
    window.close()
    app.processEvents()


def make_legacy_template(path, day):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = '调价前'
    for row, section, machine in [(3, 'A', 3), (4, 'B', 9), (5, 'C', 11)]:
        sheet.cell(row, 1, section)
        sheet.cell(row, 2, machine)
    sheet['H92'] = 123
    sheet['H93'] = 456
    discounts = workbook.create_sheet('油品优惠明细 2')
    discounts['B2'] = day
    discounts['J2'] = 0
    workbook.save(path)
    workbook.close()
    return path


def wait_for_reports(app, window):
    deadline = time.monotonic() + 10
    while window._is_processing() and time.monotonic() < deadline:
        app.processEvents()
        time.sleep(0.01)
    app.processEvents()
    assert not window._is_processing()


@pytest.mark.parametrize('simple_first', [True, False])
def test_two_templates_share_sources_but_keep_independent_outputs(
        tmp_path, monkeypatch, sources, template, shift_config, simple_first):
    app = QApplication.instance() or QApplication([])
    errors = []
    monkeypatch.setattr(QMessageBox, 'warning', lambda *args: errors.append(args[2]))
    monkeypatch.setattr(QMessageBox, 'critical', lambda *args: errors.append(args[2]))
    monkeypatch.setattr(QMessageBox, 'information', lambda *args: None)
    window = MainWindow()
    window.shift_config = shift_config
    for name in ('output_dir', 'table_dir', 'image_dir'):
        folder = tmp_path / name
        folder.mkdir()
        setattr(window, name, folder)
    # Both uploads deliberately have the same filename.
    legacy_folder = tmp_path / 'legacy_source'
    legacy_folder.mkdir()
    legacy = make_legacy_template(legacy_folder / template.name, shift_config.date.day)
    originals = {p: p.read_bytes() for p in (legacy, template)}
    uploads = [(template, 'simple_output_table'), (legacy, 'output_table')]
    for path, category in uploads if simple_first else reversed(uploads):
        window._handle_table_upload(path, category)
    assert window.output_table_path != window.simple_output_table_path
    assert window.simple_report
    assert not window.simple_options.isHidden()
    for category, path in sources.items():
        window._handle_table_upload(path, category)
    for _ in range(2):
        window._process_all_files()
        wait_for_reports(app, window)
        assert not errors
        assert window.processing_status['output_table'] == '完成'
        assert window.processing_status['simple_output_table'] == '已填写，待核对保留项'
        assert not window.pending_updates and not window.simple_pending_updates
        old = openpyxl.load_workbook(window.output_table_path)
        simple = openpyxl.load_workbook(window.simple_output_table_path)
        assert 'Sheet1' not in old.sheetnames
        assert '调价前' not in simple.sheetnames
        assert old['调价前']['H76'].value == 78
        assert old['调价前']['D3'].value == 10
        assert simple['Sheet1']['D11'].value == 164
        assert old['调价前']['H92'].value == 123
        old.close()
        simple.close()

    for category, path in [('output_table', window.output_table_path),
                           ('simple_output_table', window.simple_output_table_path)]:
        exported = tmp_path / f'{category}.xlsx'
        monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *args: (str(exported), ''))
        window.template_export_buttons[category].click()
        assert exported.read_bytes() == path.read_bytes()
    assert all(p.read_bytes() == value for p, value in originals.items())
    legacy_copy = window.output_table_path
    legacy_bytes = legacy_copy.read_bytes()
    # Replacing / clearing the simple slot never changes the legacy output.
    window._handle_table_upload(template, 'simple_output_table')
    assert legacy_copy.read_bytes() == legacy_bytes
    window._clear_category('simple_output_table')
    assert window.simple_output_table_path is None
    assert window.simple_options.isHidden()
    assert legacy_copy.read_bytes() == legacy_bytes
    assert window._validate_required_files()
    window.close()
    app.processEvents()


def test_simple_failure_does_not_block_or_modify_legacy_output(
        tmp_path, monkeypatch, sources, template, shift_config):
    app = QApplication.instance() or QApplication([])
    errors = []
    monkeypatch.setattr(QMessageBox, 'critical', lambda *args: errors.append(args[2]))
    window = MainWindow()
    window.shift_config = shift_config
    for name in ('output_dir', 'table_dir', 'image_dir'):
        folder = tmp_path / name
        folder.mkdir()
        setattr(window, name, folder)
    legacy = make_legacy_template(tmp_path / 'legacy.xlsx', shift_config.date.day)
    window._handle_table_upload(legacy, 'output_table')
    window._handle_table_upload(template, 'simple_output_table')
    before = window.simple_output_table_path.read_bytes()
    # A mismatch fails simple-report reconciliation; the legacy algorithm remains independent.
    wb = openpyxl.load_workbook(sources['油品时间统计'])
    wb.active['D4'] = 33
    wb.save(sources['油品时间统计'])
    for category, path in sources.items():
        window._handle_table_upload(path, category)
    window._process_all_files()
    wait_for_reports(app, window)
    assert not errors
    assert window.processing_status['output_table'] == '完成'
    assert window.processing_status['simple_output_table'] == '处理失败，未写入'
    assert window.simple_output_table_path.read_bytes() == before
    old = openpyxl.load_workbook(window.output_table_path)
    assert old['调价前']['D3'].value == 11
    old.close()
    assert not window.simple_pending_updates
    window.close()
    app.processEvents()


def test_template_slots_reject_wrong_type_and_allow_clearing_one(
        tmp_path, monkeypatch, template, sources, shift_config):
    app = QApplication.instance() or QApplication([])
    errors = []
    monkeypatch.setattr(QMessageBox, 'critical', lambda *args: errors.append(args[2]))
    window = MainWindow()
    for name in ('output_dir', 'table_dir', 'image_dir'):
        folder = tmp_path / name
        folder.mkdir()
        setattr(window, name, folder)
    legacy = make_legacy_template(tmp_path / 'legacy.xlsx', shift_config.date.day)
    window._handle_table_upload(legacy, 'simple_output_table')
    window._handle_table_upload(template, 'output_table')
    assert len(errors) == 2
    assert window.output_table_path is window.simple_output_table_path is None
    window._handle_table_upload(legacy, 'output_table')
    window._handle_table_upload(template, 'simple_output_table')
    for category, path in sources.items():
        if category != '油品优惠':
            window._handle_table_upload(path, category)
    assert not window._validate_required_files()  # Legacy still requires discounts.
    simple_path = window.simple_output_table_path
    window._clear_category('output_table')
    assert window.output_table_path is None
    assert window.simple_output_table_path == simple_path and simple_path.exists()
    assert window._validate_required_files()  # Simple only needs three sources.
    assert not window.simple_options.isHidden()
    window.close()
    app.processEvents()


def test_one_output_save_failure_does_not_block_the_other(
        tmp_path, monkeypatch, sources, template, shift_config):
    from src.processors.excel_updater import ExcelUpdater

    app = QApplication.instance() or QApplication([])
    errors = []
    monkeypatch.setattr(QMessageBox, 'critical', lambda *args: errors.append(args[2]))
    original_save = ExcelUpdater.save_workbook

    def fail_simple_save(updater):
        if updater.output_path.parent.name == 'simple_output_table':
            raise OSError('模拟简表保存失败')
        original_save(updater)

    monkeypatch.setattr(ExcelUpdater, 'save_workbook', fail_simple_save)
    window = MainWindow()
    window.shift_config = shift_config
    for name in ('output_dir', 'table_dir', 'image_dir'):
        folder = tmp_path / name
        folder.mkdir()
        setattr(window, name, folder)
    legacy = make_legacy_template(tmp_path / 'legacy.xlsx', shift_config.date.day)
    window._handle_table_upload(legacy, 'output_table')
    window._handle_table_upload(template, 'simple_output_table')
    before = window.simple_output_table_path.read_bytes()
    for category, path in sources.items():
        window._handle_table_upload(path, category)
    window._process_all_files()
    wait_for_reports(app, window)
    assert window.processing_status['output_table'] == '完成'
    assert window.processing_status['simple_output_table'] == '更新失败'
    assert len(errors) == 1 and '模拟简表保存失败' in errors[0]
    assert window.simple_output_table_path.read_bytes() == before
    assert not window.update_workers
    # Retry clears previous state and saves both outputs correctly.
    monkeypatch.setattr(ExcelUpdater, 'save_workbook', original_save)
    window._process_all_files()
    wait_for_reports(app, window)
    assert window.processing_status['output_table'] == '完成'
    assert window.processing_status['simple_output_table'] == '已填写，待核对保留项'
    assert not window.simple_pending_updates
    window.close()
    app.processEvents()
