"""Regression coverage using synthetic exports and optional private samples."""

import asyncio
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
from openpyxl.styles import PatternFill

from src.processors.excel_updater import ExcelUpdater
from src.processors.simple_report import SimpleReportProcessor, SIMPLE_CATEGORIES, is_simple_report
from src.processors.source_tables import read_source
from src.processors.table_processor import TableProcessor


def write_export(path, rows):
    workbook = openpyxl.Workbook()
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    return path


@pytest.fixture
def sources(tmp_path):
    rows = [
        [['油品时间统计'], ['时间：2026-10-16 8:34:40------2026-10-17 8:32:40'],
         ['站号', '机号', '油品', '加油升', '加油元', '次数'],
         [14, 3, '0#柴油 ', 30, 240, 1], [14, 9, '92#汽油 ', 30, 270, 1],
         [14, 11, '95#汽油 ', 30, 300, 1], [None, '汇总', None, 90, 810, 3]],
        [['卡型', '组', '油品', '优惠'], ['-', '卡型：充值卡', None, None],
         [None, '充值卡', '0#柴油 ', 6], [None, '充值卡', '92#汽油 ', 12],
         [None, '充值卡', '95#汽油 ', 0], [None, None, 3, 18], [3, None, 3, 18]],
        [['加油明细'], ['时间：2026-10-16 8:34:40------2026-10-17 8:32:40'],
         ['流水号', '油品', '升数', '油机金额', '结算金额', '优惠金额', '收款方式', '单价', '实收金额', 'IC卡余额'],
         [1, '0#柴油 ', 30, 240, 234, 6, '充值卡收款   ', 8, 230, 0],
         [2, '92#汽油 ', 30, 270, 258, 9, '电子卡收款   ', 9, 250, 0],
         [3, '95#汽油 ', 30, 300, 300, 0, '现金收款   ', 10, 290, 0],
         ['记录总数', 3, 90, 810, 792, 15, None]],
        [['充值明细'], ['时间：2026-10-16 8:34:40------2026-10-17 8:32:40'],
         ['卡号', '站号', '充值金额', '充值赠送', '付款方式', '原余额'],
         [1, 14, 300, 30, '现金收款   ', 100], [2, 14, 600, 0, '微信收款   ', 100],
         [3, 14, 600, 0, '支付宝收款 ', 100], [4, 14, 0, 0, None, 0],
         ['次数', None, 1500, 30, None]],
    ]
    return {category: write_export(tmp_path / f'{index}.xlsx', data)
            for index, (category, data) in enumerate(zip(SIMPLE_CATEGORIES, rows))}


@pytest.fixture
def template(tmp_path):
    path = tmp_path / 'template.xlsx'
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = 'Sheet1'
    s['A1'] = 'T-CD站数据简表'
    s.merge_cells('A1:F1')
    s['A13'] = '无卡次数'
    s['B7'] = 26
    s['B3'].fill = PatternFill('solid', fgColor='FFFF00')
    s['B6'] = '=SUM(B3:B5)'
    s['G8'] = 12345
    s['D6'] = '=SUM(D3:D5)'
    s['E6'] = '=SUM(E3:E5)'
    s2 = wb.create_sheet('Sheet2')
    s2['A4'] = '油品'
    s2['A10'] = '充值明细'
    s2['B31'] = 50
    s2['B22'] = '=SUM(B11:B21)'
    s2['C22'] = '=SUM(C11:C21)'
    wb.create_sheet('Sheet3')['A1'] = 'keep'
    wb.save(path)
    return path


@pytest.fixture
def price_change_sources(sources, tmp_path):
    before = '时间：2026-10-16 08:34:40------2026-10-16 20:00:00'
    after = '时间：2026-10-16 20:00:00------2026-10-17 08:32:40'
    full = '时间：2026-10-16 08:34:40------2026-10-17 08:32:40'
    header = ['站号', '机号', '油品', '加油升', '加油元', '次数']
    write_export(sources['油品时间统计'], [
        ['油品时间统计'], [before], header,
        [14, 3, '0#柴油', 30.01, 240.08, 1],
        [14, 9, '92#汽油', 3.01, 27.09, 1],
        [14, 11, '95#汽油', 6.01, 60.10, 1],
        ['合计', None, None, 39.03, 327.27, 3],
    ])
    sources['油品时间统计(调价后)'] = write_export(tmp_path / 'after.xlsx', [
        ['油品时间统计'], [after], header,
        [14, 3, '0#柴油', 30.01, 270.09, 1],
        [14, 9, '92#汽油', 3.01, 30.10, 1],
        [14, 11, '95#汽油', 6.01, 66.11, 1],
        ['合计', None, None, 39.03, 366.30, 3],
    ])
    write_export(sources['加油明细'], [
        ['加油明细'], [full],
        ['油品', '单价', '升数', '油机金额', '结算金额', '优惠金额', '收款方式', 'IC卡余额', '加油完成时间', '实收金额'],
        ['0#柴油', 8, 30.01, 240.08, 240.08, 0, '现金收款', 0, '2026-10-16 19:00:00', 240.08],
        ['0#柴油', 9, 30.01, 270.09, 270.09, 0, '现金收款', 0, '2026-10-16 20:00:00', 270.09],
        ['92#汽油', 9, 3.01, 27.09, 27.09, 0, '现金收款', 0, '2026-10-16 19:00:00', 27.09],
        ['92#汽油', 10, 3.01, 30.10, 30.10, 0, '现金收款', 0, '2026-10-16 20:00:00', 30.10],
        ['95#汽油', 10, 6.01, 60.10, 60.10, 0, '现金收款', 0, '2026-10-16 19:00:00', 60.10],
        ['95#汽油', 11, 6.01, 66.11, 66.11, 0, '现金收款', 0, '2026-10-16 20:00:00', 66.11],
    ])
    sources.pop('油品优惠')
    return sources


def test_price_change_combines_all_three_metrics(price_change_sources, shift_config):
    result = SimpleReportProcessor(shift_config).build_updates(price_change_sources)
    first = result['updates'][0]['cells']
    for cell, expected in {
        'B3': 20, 'D3': 170, 'E3': 1,
        'B4': 2, 'D4': 19, 'E4': 1,
        'B5': 4, 'D5': 42, 'E5': 1,
    }.items():
        assert first[cell] == expected, cell
    assert result['processed_data']['原始升数'] == 78.06
    assert result['processed_data']['原始油机金额'] == 693.57
    assert result['processed_data']['原始次数'] == 6
    assert '调价后' in result['processed_data']['时间统计']


def test_same_price_splits_preserve_original_algorithm(sources, tmp_path, shift_config):
    header = ['站号', '机号', '油品', '加油升', '加油元', '次数']
    write_export(sources['油品时间统计'], [header, [14, 9, '92#汽油', 123.65, 1075.73, 1]])
    sources['油品时间统计(调价后)'] = write_export(tmp_path / 'after.xlsx', [
        header, [14, 9, '92#汽油', 315.02, 2739.85, 1]])
    write_export(sources['加油明细'], [
        ['油品', '单价', '升数', '油机金额', '结算金额', '优惠金额', '收款方式', 'IC卡余额', '实收金额'],
        ['92#汽油', 8.7, 123.65, 1075.73, 1075.73, 0, '现金收款', 0, 1075.73],
        ['92#汽油', 8.7, 315.02, 2739.85, 2739.85, 0, '现金收款', 0, 2739.85]])
    sources.pop('油品优惠')
    first = SimpleReportProcessor(shift_config).build_updates(sources)['updates'][0]['cells']
    assert (first['B4'], first['D4'], first['E4']) == (146.22, 1272.11, 1)


def test_actual_price_change_requires_transaction_times(price_change_sources, shift_config):
    wb = openpyxl.load_workbook(price_change_sources['加油明细'])
    wb.active['I4'] = 'bad-time'
    wb.save(price_change_sources['加油明细'])
    with pytest.raises(ValueError, match='加油完成时间'):
        SimpleReportProcessor(shift_config).build_updates(price_change_sources)


@pytest.mark.parametrize('cell,value,message', [
    ('D4', 31.01, '升数不一致'),
    ('E4', 271.09, '油机金额不一致'),
    ('F4', 2, '次数不一致'),
    ('A4', 15, '站号不一致'),
    ('A2', '时间：2026-10-16 20:00:00------2026-10-17 07:32:40', '时段不一致'),
    ('C4', '98#汽油', '未支持的油品'),
    ('E3', '缺失金额列', '找不到表头'),
])
def test_invalid_after_statistics_fail(price_change_sources, shift_config, cell, value, message):
    wb = openpyxl.load_workbook(price_change_sources['油品时间统计(调价后)'])
    wb.active[cell] = value
    wb.save(price_change_sources['油品时间统计(调价后)'])
    with pytest.raises(ValueError, match=message):
        SimpleReportProcessor(shift_config).build_updates(price_change_sources)


def test_four_sources_reconcile_and_map(sources, shift_config):
    result = SimpleReportProcessor(shift_config).build_updates(sources)
    first, second = [update['cells'] for update in result['updates']]
    assert [first[f'A{r}'] for r in (3, 4, 5)] == ['0#柴油', '92#汽油', '95#汽油']
    assert [first[f'B{r}'] for r in (3, 4, 5)] == [10, 10, 10]
    assert [first[f'D{r}'] for r in (3, 4, 5)] == [80, 90, 100]
    assert first['E4'] == 0
    assert first['D11'] == 164
    assert first['D12'] == 86
    assert second['D31'] == 4  # Includes the coupon, absent from 优惠金额.
    assert second['D34'] == 2
    assert first['B12'] == 400
    assert first['D9'] == 510  # Accepted payments include recharge gifts.
    assert first['B9'] == 1.33  # Includes the account-opening record.
    assert first['D13'] == 96.67  # No-card uses actual receipts (290 / 3).
    assert first['B13'] == 0  # Scaled counts round to a whole number.
    assert result['processed_data']['折算'] == '金额/升数 ÷ 3；次数 ÷ 3'
    assert 'A13' not in second  # Unspecified recharge-denomination cells stay untouched.
    assert first['B8'] is first['D8'] is None
    assert second['D7'] == '=IF(C7=0,0,B7/C7)'
    assert first['F6'] == '=IF(E6=0,0,D6/E6)'


@pytest.mark.parametrize('divisor,count_divisor', [(1, 1), (3, 1), (3, 3)])
def test_scaling(sources, shift_config, divisor, count_divisor):
    result = SimpleReportProcessor(shift_config, divisor, count_divisor).build_updates(sources)
    first, second = [u['cells'] for u in result['updates']]
    assert first['B3'] == 30 / divisor
    assert first['E3'] == (1 if count_divisor == 1 else 0)
    assert first['D9'] == 1530 / divisor


@pytest.mark.parametrize('divisor', [0, -1, float('inf'), float('nan')])
def test_invalid_scaling(shift_config, divisor):
    with pytest.raises(ValueError):
        SimpleReportProcessor(shift_config, divisor)


@pytest.mark.parametrize('category,cell,value,message', [
    ('油品时间统计', 'D4', 31, '升数不一致'),
    ('油品时间统计', 'E4', 241, '油机金额不一致'),
    ('油品时间统计', 'F4', 2, '次数不一致'),
    ('油品优惠', 'D3', 7, '优惠不一致'),
    ('油品时间统计', 'C4', '98#汽油', '未支持的油品'),
    ('加油明细', 'E4', '损坏', '无效数值'),
    ('充值明细', 'A2', '时间：2026-10-15 8:34:40------2026-10-16 8:32:40', '时段不一致'),
    ('充值明细', 'B4', 15, '站号不一致'),
])
def test_invalid_sources_fail_before_writes(sources, shift_config, category, cell, value, message):
    wb = openpyxl.load_workbook(sources[category])
    wb.active[cell] = value
    wb.save(sources[category])
    with pytest.raises(ValueError, match=message):
        SimpleReportProcessor(shift_config).build_updates(sources)


def test_idempotent_updates_preserve_template(sources, template, shift_config):
    result = SimpleReportProcessor(shift_config).build_updates(sources)
    updater = ExcelUpdater(template)
    assert is_simple_report(updater.workbook)
    before_style = updater.workbook['Sheet1']['B3'].style_id
    for _ in range(2):
        updater.apply_updates(result['updates'])
    updater.save_workbook()
    updater.close_workbook()
    wb = openpyxl.load_workbook(template)
    assert wb['Sheet1']['B3'].value == 10
    assert wb['Sheet1']['B3'].style_id == before_style
    assert wb['Sheet1']['E3'].number_format == '0'
    assert wb['Sheet2']['B11'].number_format == 'General'
    assert wb['Sheet1']['B7'].value == 26
    assert wb['Sheet1']['G8'].value == 12345
    assert wb['Sheet2']['B31'].value == 50
    assert wb['Sheet3']['A1'].value == 'keep'
    assert 'A1:F1' in wb['Sheet1'].merged_cells


def test_invalid_batch_does_not_partially_write(template):
    updater = ExcelUpdater(template)
    with pytest.raises(ValueError):
        updater.apply_updates([{'sheet': 'Sheet1', 'cells': {'B7': 0}},
                               {'sheet': 'Sheet1', 'cells': {'B1': 9}}])
    assert updater.workbook['Sheet1']['B7'].value == 26


def test_legacy_whitespace_and_grouped_discounts(sources, shift_config):
    processor = TableProcessor(shift_config)
    result = asyncio.run(processor.process_table(sources['加油明细'], '加油明细'))
    assert result['processed_data'] == {'customer_discount': 78, 'electric_discount': 86}
    assert result['updates'][0]['updates'][0]['row'] == 76
    result = asyncio.run(processor.process_table(sources['油品优惠'], '油品优惠'))
    assert result['processed_data'] == {'gasoline_discount': 4, 'diesel_discount': 2}
    result = asyncio.run(processor.process_table(sources['充值明细'], '充值明细'))
    assert result['processed_data'] == {'online_recharge': 400, 'cash_recharge': 110}


def test_missing_header(tmp_path):
    path = write_export(tmp_path / 'bad.xlsx', [['wrong'], [1]])
    with pytest.raises(ValueError, match='找不到表头'):
        read_source(path, ['油品'])


def test_zero_activity_clears_old_values(sources, shift_config):
    for category, path in sources.items():
        wb = openpyxl.load_workbook(path)
        first_data = 2 if category == '油品优惠' else 4
        wb.active.delete_rows(first_data, wb.active.max_row)
        wb.save(path)
    result = SimpleReportProcessor(shift_config).build_updates(sources)
    first, second = [u['cells'] for u in result['updates']]
    assert first['B3'] == first['D3'] == first['E3'] == first['B12'] == 0
    assert 'F10' not in first
    assert first['F13'] == '=IF(SUM(D4:D5)=0,0,D13/SUM(D4:D5))'
    assert 'A11' not in second
    assert second['D31'] == second['D34'] == 0


def test_online_recharge_includes_gifts_refunds_and_payment_labels(sources, shift_config):
    processor = SimpleReportProcessor(shift_config)
    recharge = pd.DataFrame([
        [100, 18144, 240, '微信收款'],
        [100, 3700, 0, '支付宝收款'],
        [100, -300, -30, '微信退款'],
        [100, 600, 30, '线上支付宝收款'],
        [100, 999, 99, '现金收款'],
    ], columns=['原余额', '充值金额', '充值赠送', '付款方式'])
    cells = {}
    processor._recharge_cells(recharge, cells)
    assert cells['B12'] == 7461.33


def test_signed_recharge_and_new_account(sources, shift_config):
    path = sources['充值明细']
    wb = openpyxl.load_workbook(path)
    wb.active['C4'] = -300
    wb.active['F5'] = 0  # A new account with an actual payment.
    wb.active.append([5, 14, 123, 0, None, 100])  # Reversal: count only.
    wb.save(path)
    first = SimpleReportProcessor(shift_config).build_updates(sources)['updates'][0]['cells']
    assert first['D9'] == 110  # (-300 + 30 + 600) / 3
    assert first['D8'] == 200
    assert first['B8'] == 0.33
    assert first['B9'] == 1.67  # All five records, including opening/reversal.
    assert first['B12'] == 400


def test_card_methods_round_separately(sources, shift_config):
    wb = openpyxl.load_workbook(sources['加油明细'])
    wb.active['E4'] = 234.02
    wb.active['E5'] = 258.02
    wb.save(sources['加油明细'])
    # Discount source is optional; keep independently validating volumes/money/counts.
    paths = {k: v for k, v in sources.items() if k != '油品优惠'}
    result = SimpleReportProcessor(shift_config).build_updates(paths)
    first, second = [u['cells'] for u in result['updates']]
    assert first['D11'] == 164.02  # 78.01 + 86.01, not combined-total 164.01.
    assert first['D12'] == 86.01  # Settlement differs from actual receipts (250 / 3).
    assert 'D31' not in second and 'D34' not in second
    assert '未提供优惠表' in result['processed_data']['交叉校验']


@pytest.mark.parametrize('machines,litres,amount', [([1, 2], 0.68, 5.92), ([1, 1], 0.67, 5.83)])
def test_machine_rounding_before_fuel_total(sources, shift_config, machines, litres, amount):
    write_export(sources['油品时间统计'], [
        ['机号', '油品', '加油升', '加油元', '次数'],
        [machines[0], '92#汽油', 1.01, 8.79, 1], [machines[1], '92#汽油', 1.01, 8.79, 1],
    ])
    write_export(sources['加油明细'], [
        ['流水号', '油品', '升数', '油机金额', '结算金额', '优惠金额', '收款方式', '单价', '实收金额', 'IC卡余额'],
        [1, '92#汽油', 1.01, 8.79, 8.79, 0, '现金收款', 8.7, 8.79, 0],
        [2, '92#汽油', 1.01, 8.79, 8.79, 0, '现金收款', 8.7, 8.79, 0],
    ])
    sources.pop('油品优惠')
    first = SimpleReportProcessor(shift_config).build_updates(sources)['updates'][0]['cells']
    assert first['B4'] == litres
    assert first['D4'] == amount
    assert first['E4'] == 1


def test_no_card_uses_actual_receipts_ignores_balance_and_excludes_tank_returns(sources, shift_config):
    write_export(sources['油品时间统计'], [
        ['机号', '油品', '加油升', '加油元', '次数'],
        [1, '0#柴油', 10, 80, 1], [2, '92#汽油', 40, 360, 4], [3, '95#汽油', 30, 300, 3],
    ])
    write_export(sources['加油明细'], [
        ['流水号', '油品', '升数', '油机金额', '结算金额', '优惠金额', '收款方式', '单价', '实收金额', 'IC卡余额', '卡号'],
        [1, '0#柴油', 10, 80, 80, 0, '现金收款', 8, 70, 0, 0],  # Diesel excluded.
        [2, '92#汽油', 10, 90, 0, 0, '现金收款', 9, 0, 0, 0],  # Zero settlement excluded.
        [3, '92#汽油', 10, 90, 85, 5, '现金收款', 9, 80, 0, 0],  # Discount excluded.
        [4, '95#汽油', 10, 100, 100, 0, '充值卡收款', 10, 93, 200, 0],
        [5, '92#汽油', 10, 90, 90, 0, '现金收款', 9, 90, 10, 0],  # Balance does not affect eligibility.
        [6, '95#汽油', 10, 100, 100, 0, '回罐收款  ', 10, 100, 0, 9321],
        [7, '92#汽油', 10, 90, 90, 0, '现金收款', 9, 90, 0, 123],  # Card number is not the filter.
        [8, '95#汽油', 10, 100, 100, 0, '回灌收款', 10, 100, 0, 0],
    ])
    sources.pop('油品优惠')
    first = SimpleReportProcessor(shift_config, count_divisor=1).build_updates(sources)['updates'][0]['cells']
    assert first['B13'] == 3
    assert first['D13'] == 91  # Actual receipts (93 + 90 + 90) / 3.


@pytest.mark.parametrize('balance', [None, 'bad', 'NaN', 'Infinity'])
def test_no_card_does_not_validate_card_balance(sources, shift_config, balance):
    wb = openpyxl.load_workbook(sources['加油明细'])
    wb.active['J6'] = balance
    wb.save(sources['加油明细'])
    first = SimpleReportProcessor(shift_config).build_updates(sources)['updates'][0]['cells']
    assert first['D13'] == 96.67


def test_card_balance_header_is_optional(sources, shift_config):
    wb = openpyxl.load_workbook(sources['加油明细'])
    wb.active.delete_cols(10)
    wb.save(sources['加油明细'])
    first = SimpleReportProcessor(shift_config).build_updates(sources)['updates'][0]['cells']
    assert first['D13'] == 96.67


@pytest.mark.parametrize('missing', [True, False])
def test_actual_receipts_are_required_for_no_card(sources, shift_config, missing):
    workbook = openpyxl.load_workbook(sources['加油明细'])
    if missing:
        workbook.active.delete_cols(9)
    else:
        workbook.active['I6'] = 'unused'
    workbook.save(sources['加油明细'])
    with pytest.raises(ValueError, match='实收金额'):
        SimpleReportProcessor(shift_config).build_updates(sources)


def test_card_consumption_keeps_settlement_when_actual_receipts_differ(sources, shift_config):
    workbook = openpyxl.load_workbook(sources['加油明细'])
    workbook.active['I4'] = 'unused'
    workbook.active['I5'] = 'unused'
    workbook.save(sources['加油明细'])
    first = SimpleReportProcessor(shift_config).build_updates(sources)['updates'][0]['cells']
    assert (first['D11'], first['D12'], first['D13']) == (164, 86, 96.67)


def test_multiple_fuel_prices_require_split_period(sources, shift_config):
    wb = openpyxl.load_workbook(sources['加油明细'])
    wb.active.append([4, '92#汽油', 0, 0, 0, 0, '现金收款', 10, 0, 0])
    wb.save(sources['加油明细'])
    wb = openpyxl.load_workbook(sources['油品时间统计'])
    wb.active['F5'] = 2
    wb.save(sources['油品时间统计'])
    with pytest.raises(ValueError, match='多个单价'):
        SimpleReportProcessor(shift_config).build_updates(sources)


@pytest.mark.parametrize('value,expected', [('608377.84', 202792.61), ('0', 0), ('', None)])
def test_only_requested_balance_cell(sources, shift_config, value, expected):
    first = SimpleReportProcessor(shift_config, closing_balance=value).build_updates(sources)['updates'][0]['cells']
    assert first.get('G7') == expected
    assert ('G7' in first) == (expected is not None)
    assert all(cell not in first for cell in ['G8', 'H8', 'I8', 'J8', 'B10', 'D10'])


@pytest.mark.parametrize('value', ['NaN', 'Infinity', 'bad'])
def test_invalid_balance(shift_config, value):
    with pytest.raises(ValueError, match='期末余额'):
        SimpleReportProcessor(shift_config, closing_balance=value)


def test_verified_oct20_sample(tmp_path, shift_config):
    """Keep the verified baseline, including the restored actual-receipts rule."""
    base = Path(__file__).resolve().parents[1] / '.context/attachments'
    paths = {'油品时间统计': base / 'H08NC1/1(5).xlsx',
             '加油明细': base / 'cCDPeT/3(5).xlsx',
             '充值明细': base / 'mKlowY/4(5).xlsx'}
    expected_path = base / '0LPzc8/T-CD简20(1).xlsx'
    if not expected_path.exists() or not all(p.exists() for p in paths.values()):
        pytest.skip('Private verified attachments are not available')
    result = SimpleReportProcessor(shift_config).build_updates(paths)
    first = result['updates'][0]['cells']
    expected = openpyxl.load_workbook(expected_path, data_only=True)
    addresses = [f'{col}{row}' for row in (3, 4, 5) for col in ('B', 'D', 'E')]
    addresses += ['B8', 'D8', 'B9', 'D9', 'D11', 'B12', 'D12', 'B13']
    for address in addresses:
        assert first[address] == expected['Sheet1'][address].value, address
    assert expected['Sheet1']['D13'].value == 1550.99
    assert first['D13'] == 1550.99
    addresses.append('D13')
    expected.close()
    # Exercise save/reopen/idempotence and preserve every cell outside the update batch.
    original = expected_path.read_bytes()
    output = tmp_path / 'verified.xlsx'
    output.write_bytes(original)
    before = openpyxl.load_workbook(expected_path)
    updater = ExcelUpdater(output)
    updater.apply_updates(result['updates'])
    updater.apply_updates(result['updates'])
    updater.save_workbook()
    updater.close_workbook()
    after = openpyxl.load_workbook(output)
    changed = {u['sheet']: set(u['cells']) for u in result['updates']}
    for sheet in before:
        for row in sheet:
            for cell in row:
                if cell.coordinate not in changed.get(sheet.title, set()):
                    assert after[sheet.title][cell.coordinate].value == cell.value
    for address in addresses:
        assert after['Sheet1'][address].value == first[address]
    assert expected_path.read_bytes() == original
    before.close()
    after.close()
